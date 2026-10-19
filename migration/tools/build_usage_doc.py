"""Build the Chinese user guide for 抽屉 Drawer (docs/抽屉使用说明.docx).

Usage:
    "<bundled-python>" migration/tools/build_usage_doc.py

The screenshots under docs/images come from `capture_usage_shots.py` (seeded
demo data, no personal paths). Tweak the copy below, rerun, then export a PDF
with the LibreOffice Kit if needed.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO = Path(__file__).resolve().parents[2]
SHOTS = REPO / "docs" / "images"
OUT = REPO / "docs" / "抽屉使用说明.docx"

TEXT_FONT = "Microsoft YaHei"
TEXT_FONT_EA = "微软雅黑"
MONO_FONT = "Consolas"
INK = RGBColor(0x1F, 0x23, 0x28)
MUTED = RGBColor(0x5F, 0x63, 0x68)
BLUE = RGBColor(0x0B, 0x57, 0xD0)
HEADER_FILL = "DCE7FB"
CALLOUT_FILL = "EAF1FE"

_figure = 0


def _set_run_font(run, mono: bool = False) -> None:
    name = MONO_FONT if mono else TEXT_FONT
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), MONO_FONT if mono else TEXT_FONT_EA)


def _style_font(style, size: float, bold: bool = False, color: RGBColor = INK) -> None:
    style.font.name = TEXT_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), TEXT_FONT)
    rfonts.set(qn("w:hAnsi"), TEXT_FONT)
    rfonts.set(qn("w:eastAsia"), TEXT_FONT_EA)


def add_rich(paragraph, text: str, size: float | None = None, bold: bool = False,
             color: RGBColor | None = None) -> None:
    """Add text with **bold** and `code` markup."""
    for token in re.split(r"(\*\*.*?\*\*|`.*?`)", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            _set_run_font(run)
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, mono=True)
            run.font.size = Pt((size or 10.5) - 0.5)
            run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        else:
            run = paragraph.add_run(token)
            _set_run_font(run)
            run.bold = bold
            if color is not None:
                run.font.color.rgb = color
        if size is not None:
            run.font.size = Pt(size)
        if bold:
            run.bold = True


def body(doc, text: str, space_after: float = 5, indent: float = 0) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(0)
    if indent:
        paragraph.paragraph_format.left_indent = Cm(indent)
    add_rich(paragraph, text)


def bullet(doc, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.left_indent = Cm(0.7 + level * 0.6)
    add_rich(paragraph, text)


def numbered(doc, items: list[str]) -> None:
    for index, text in enumerate(items, start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.left_indent = Cm(0.85)
        paragraph.paragraph_format.first_line_indent = Cm(-0.55)
        add_rich(paragraph, f"{index}. {text}")


def h1(doc, text: str) -> None:
    heading = doc.add_heading(text, level=1)
    heading.paragraph_format.space_before = Pt(16)
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.keep_with_next = True
    for run in heading.runs:
        _set_run_font(run)
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = INK
    # Thin blue rule under the section title.
    ppr = heading._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "0B57D0")
    borders.append(bottom)
    ppr.append(borders)


def h2(doc, text: str) -> None:
    heading = doc.add_heading(text, level=2)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(3)
    heading.paragraph_format.keep_with_next = True
    for run in heading.runs:
        _set_run_font(run)
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = INK


def shade(cell, hex_fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def set_cell_margins(table, top: int = 60, bottom: int = 60, left: int = 90, right: int = 90) -> None:
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        mar.append(element)
    tblpr.append(mar)


def table(doc, headers: list[str], rows: list[list[str]], widths: list[float] | None = None,
          font_size: float = 9.0) -> None:
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_cell_margins(tbl)
    for index, text in enumerate(headers):
        cell = tbl.rows[0].cells[index]
        shade(cell, HEADER_FILL)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        add_rich(paragraph, text, size=font_size, bold=True)
    for row in rows:
        cells = tbl.add_row().cells
        for index, text in enumerate(row):
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            add_rich(paragraph, text, size=font_size)
    if widths:
        for index, width in enumerate(widths):
            for row in tbl.rows:
                row.cells[index].width = Cm(width)
    # Repeat the header row when a long table continues on the next page, and
    # never split a single row across pages.
    header_trpr = tbl.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_trpr.append(repeat)
    for row in tbl.rows:
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)


def callout(doc, title: str, text: str) -> None:
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    set_cell_margins(tbl, top=110, bottom=110, left=160, right=160)
    cell = tbl.rows[0].cells[0]
    shade(cell, CALLOUT_FILL)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(f"{title}  ")
    _set_run_font(run)
    run.bold = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = BLUE
    add_rich(paragraph, text, size=9.5)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)


def figure(doc, filename: str, caption: str, width_cm: float = 15.6) -> None:
    global _figure
    _figure += 1
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run()
    run.add_picture(str(SHOTS / filename), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    add_rich(cap, f"图 {_figure} · {caption}", size=8.5, color=MUTED)


def add_page_number_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    _set_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    inner = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "5F6368")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    rpr.append(color)
    rpr.append(size)
    inner.append(rpr)
    text = OxmlElement("w:t")
    text.text = "1"
    inner.append(text)
    fld.append(inner)
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    _set_run_font(run)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def header(section) -> None:
    paragraph = section.header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("抽屉 Drawer · 使用说明")
    _set_run_font(run)
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED


def cover(doc) -> None:
    for _ in range(4):
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("抽屉 Drawer")
    _set_run_font(run)
    run.font.size = Pt(34)
    run.bold = True
    run.font.color.rgb = INK

    subtitle = doc.add_paragraph()
    run = subtitle.add_run("使用说明")
    _set_run_font(run)
    run.font.size = Pt(22)
    run.font.color.rgb = BLUE
    subtitle.paragraph_format.space_after = Pt(10)

    tagline = doc.add_paragraph()
    run = tagline.add_run("把常用网站、文件和文件夹，收进一个随时能打开的抽屉。")
    _set_run_font(run)
    run.font.size = Pt(12)
    run.font.color.rgb = MUTED
    tagline.paragraph_format.space_after = Pt(26)

    meta = doc.add_paragraph()
    add_rich(meta, "版本 1.0 · 2026-10-18 · 适用于 Windows 10 / 11（64 位）", size=10, color=MUTED)
    meta.paragraph_format.space_after = Pt(6)

    rule = doc.add_paragraph()
    ppr = rule._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "0B57D0")
    borders.append(bottom)
    ppr.append(borders)

    abstract = doc.add_paragraph()
    abstract.paragraph_format.space_before = Pt(18)
    abstract.paragraph_format.line_spacing = 1.5
    add_rich(
        abstract,
        "抽屉是一个**完全本地运行**的小工具：把常用的网站、文件和文件夹收藏起来，"
        "单击就能在右侧预览文件夹的思维导图；按一个快捷键，还能在任意软件上方直接弹出"
        "搜索面板。它不注册账号、不上传数据，除了可选的 AI 分类建议之外不需要联网。",
        size=10.5,
    )

    hint = doc.add_paragraph()
    hint.paragraph_format.space_before = Pt(10)
    add_rich(
        hint,
        "本手册按「安装 → 主界面 → 添加资源 → 思维导图 → 快捷键面板 → 书签同步 → "
        "右键菜单 → AI → 设置/数据/常见问题」的顺序介绍全部功能。文中的截图使用示例数据，"
        "你的界面内容会以自己的资源为准。",
        size=9.5,
        color=MUTED,
    )

    doc.add_page_break()


def build() -> Document:
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    header(section)
    add_page_number_footer(section)
    section.different_first_page_header_footer = True

    _style_font(doc.styles["Normal"], 10.5)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.4
    normal.space_after = Pt(5)

    doc.core_properties.title = "抽屉 Drawer 使用说明"
    doc.core_properties.author = "抽屉 Drawer"
    doc.core_properties.subject = "抽屉 Drawer 用户手册"

    cover(doc)

    # ------------------------------------------------------------------ 1
    h1(doc, "1. 安装与启动")
    h2(doc, "1.1 系统要求")
    bullet(doc, "Windows 10（1809 及以上）或 Windows 11，**64 位**系统。")
    bullet(doc, "约 50 MB 磁盘空间。")
    bullet(doc, "**不需要**安装 Python、.NET 或其它运行库——程序已经自带。")
    bullet(doc, "第一次运行会在程序目录写入数据文件，因此要放在**可写**的文件夹里。")

    h2(doc, "1.2 安装（绿色免安装）")
    numbered(
        doc,
        [
            "把收到的压缩包**整包解压**成一个文件夹，例如 `D:\\抽屉`。不要只单独拷贝 "
            "`resource_hub.exe`——它还需要同目录下的 `data\\`、`backend\\`、"
            "`flutter_windows.dll` 等文件。",
            "双击 `resource_hub.exe` 启动。",
            "如果 Windows 弹出蓝色的「已保护你的电脑」（未签名程序的 SmartScreen 提示），"
            "点「**更多信息**」→「**仍要运行**」。如果杀毒软件把 "
            "`backend\\resource_hub_backend.exe` 误报为风险，把它加入信任即可。",
            "首次启动会自动生成几条示例资源，可以直接删除。",
        ],
    )
    callout(
        doc,
        "建议位置",
        "不要放在 `C:\\Program Files` 等没有写权限的目录，否则收藏和设置会保存失败。"
        "放在 D 盘、桌面上的普通文件夹都可以；以后整个文件夹移动位置也没关系，数据会跟着走。",
    )

    h2(doc, "1.3 文件夹里有什么")
    table(
        doc,
        ["文件 / 目录", "作用"],
        [
            ["resource_hub.exe", "主程序，双击启动"],
            ["backend\\", "内置的数据服务，请勿删除或单独移动"],
            ["data\\、flutter_windows.dll 等", "界面运行库，请勿删除"],
            ["bookmark-extension\\", "浏览器扩展（可选，用于双向书签同步）"],
            ["sites.json / folders.json", "你的网站 / 本地资源数据，首次运行自动生成"],
            ["settings.json", "设置（语言、快捷键、开关等）"],
            [".resourcehub-meta.json", "书签同步与排序状态，自动维护，不用手改"],
            [".resourcehub-ai-key", "配置过 AI 后保存的加密密钥（没有配置就没有这个文件）"],
        ],
        widths=[5.2, 11.0],
    )

    # ------------------------------------------------------------------ 2
    h1(doc, "2. 主界面速览")
    body(doc, "主窗口分为四块：顶部标题、工具栏（搜索 / 筛选 / 操作）、中间的列表或网格、"
              "以及从右侧滑出的详情面板。")
    figure(doc, "main-window.png", "主界面（示例数据）")

    h2(doc, "2.1 顶栏")
    bullet(doc, "左侧「**已保存资源 (N)**」显示当前分类下的资源数量。")
    bullet(doc, "右侧三个控件：**网站 / 资源** 切换、**列表 / 网格** 切换、蓝色的 **+** 新增按钮。")

    h2(doc, "2.2 搜索、筛选与批量操作")
    bullet(doc, "**搜索框**：输入即筛选，匹配名称、路径 / 网址、类型、分类和备注；清空搜索后恢复完整列表。")
    bullet(doc, "**分类筛选**：下拉选择「全部分类」「未分类」或某个具体分类。")
    bullet(doc, "**只看收藏**（星形按钮）：只显示已加星的资源。")
    bullet(doc, "**操作按钮**：打开、收藏 / 取消收藏、检测路径、编辑、删除。"
                 "未选中资源时这些按钮是灰色的；检测路径不需要选中，会检查当前模式下的全部资源。")
    bullet(doc, "**更多菜单**：按添加时间排序、同步浏览器书签、AI 整理建议、导出、导入。")
    bullet(doc, "**齿轮**：打开设置。")

    h2(doc, "2.3 列表与网格")
    bullet(doc, "**单击**一行 / 一张卡片：选中它，并在右侧打开详情面板。")
    bullet(doc, "**双击**：直接打开资源（网站用浏览器打开，文件夹用资源管理器打开，文件用默认程序打开）。")
    bullet(doc, "**右键**：打开 / 打开所在位置 / 编辑 / 收藏 / 复制路径（网站则是复制网址）/ 删除。")
    bullet(doc, "收藏过的资源会**自动置顶**显示；列表视图显示完整信息，网格视图适合快速浏览。")

    h2(doc, "2.4 右侧详情面板")
    bullet(doc, "文件夹资源：显示**思维导图**，可以一层层展开预览（详见第 4 章）。")
    bullet(doc, "文件或网站资源：显示分类、加入时间、备注和路径 / 网址，可直接复制。")
    bullet(doc, "面板右上角的 **×** 或按 **Esc** 关闭。")

    h2(doc, "2.5 底部状态栏")
    body(doc, "每次操作的结果（已保存、已删除、发现失效路径、AI 建议数量等）都会在左下角提示；"
              "右侧显示当前列表的条数。")

    # ------------------------------------------------------------------ 3
    h1(doc, "3. 添加与管理资源")
    h2(doc, "3.1 新增网站")
    numbered(
        doc,
        [
            "点击顶部的「**网站**」切换到网站模式。",
            "点击蓝色的 **+** 按钮。",
            "填写名称、网址、分类和备注（加入时间自动记录），点「保存链接」。",
        ],
    )
    body(doc, "网址会做基本校验，重复的网址会被拒绝并提示。")

    h2(doc, "3.2 新增文件 / 文件夹")
    numbered(
        doc,
        [
            "切换到「**资源**」模式，点击 **+**。",
            "填写名称；在「本地路径」里可以直接粘贴路径，也可以点右侧的文件夹 / 文件按钮选择。"
            "工具会自动识别你选的是文件夹还是文件。",
            "选择分类、填写备注，点「保存资源」。",
        ],
    )

    h2(doc, "3.3 编辑、收藏与删除")
    bullet(doc, "**编辑**：选中资源后点铅笔按钮，或右键「编辑」；改完点「保存修改」。")
    bullet(doc, "**收藏**：点星形按钮或右键「收藏」，收藏后自动置顶。")
    bullet(doc, "**删除**：点垃圾桶按钮或右键「删除」，在确认框里点「删除」。")
    body(
        doc,
        "如果删除的网站同时存在于 Chrome / Edge 书签里，确认框会多出一个勾选项"
        "「同时从浏览器移除（N 处书签收藏）」，并可以「记住我的选择」——"
        "对应的删除策略在设置里也能改（每次询问 / 同时移除书签 / 仅删本地）。",
    )

    h2(doc, "3.4 检测路径")
    body(doc, "点盾牌按钮会检查当前模式下所有本地资源的路径。移动或删除了文件夹后，"
              "状态栏会提示有几条路径失效；对失效的资源重新编辑路径，或删除后重新添加即可。")

    h2(doc, "3.5 排序")
    bullet(doc, "**拖动排序**：列表视图按住行首的拖动图标，网格视图直接按住卡片拖动，松手后自动保存；"
                 "顶部会出现「顺序已更新」，可以点「撤销」恢复。")
    bullet(doc, "搜索状态下不能拖动排序，清空搜索即可。")
    bullet(doc, "**Alt + ↑ / ↓**：把当前选中的资源上移 / 下移一位。")
    bullet(doc, "**更多菜单 → 按添加时间排序**：按加入时间从早到晚重新排列。")

    h2(doc, "3.6 搜索与分类筛选")
    body(doc, "搜索会同时匹配名称、路径 / 网址、类型（文件夹 / 文件 / 网站）、分类和备注。"
              "搜索和筛选可以叠加；分类下拉里的分类来自现有资源、浏览器书签的文件夹名以及 AI 建议。")

    h2(doc, "3.7 导出与导入")
    bullet(doc, "**更多菜单 → 导出**：把当前模式（网站或资源）保存成一个 JSON 备份文件。")
    bullet(doc, "**更多菜单 → 导入**：选择 JSON 备份文件，**只新增**不覆盖，重复项会自动跳过并提示数量。")
    bullet(doc, "备份文件与旧的 Python 版本格式完全兼容，可以互相导入导出。")

    # ------------------------------------------------------------------ 4
    h1(doc, "4. 文件夹详情：思维导图")
    body(doc, "单击一个文件夹资源，右侧会滑出它的思维导图。根节点是文件夹本身，"
              "点节点右侧的小圆箭头可以展开下一层，收起后再点一次即可折叠。")
    figure(doc, "folder-mindmap.png", "文件夹思维导图（示例数据）", width_cm=8.8)

    h2(doc, "4.1 节点操作")
    bullet(doc, "**单击节点**：选中（右侧出现展开箭头等操作）。")
    bullet(doc, "**双击节点**：打开——文件夹用资源管理器，文件用默认程序。")
    bullet(doc, "**右键节点**：打开 / 打开所在位置 / 复制路径 / 添加为资源（把子文件夹或文件收进资源库）。")
    bullet(doc, "**小圆箭头**：展开 / 收起下一层。文件夹内容是一层一层按需读取的，大文件夹也不会卡。")
    bullet(doc, "文件夹里的条目很多时，节点下方会显示「还有 N 项未显示」（单层最多显示 200 项）。")

    h2(doc, "4.2 缩放、平移与适应视图")
    table(
        doc,
        ["操作", "效果"],
        [
            ["按住空白处或节点拖动", "平移视图（鼠标变成抓手）"],
            ["滚轮", "上下平移；按住 Shift 滚轮则左右平移"],
            ["Ctrl + 滚轮", "以鼠标位置为中心缩放（40%～200%）"],
            ["工具栏 − / 百分比 / +", "缩小、查看当前比例、放大"],
            ["工具栏「适应视图」或双击空白处", "自动缩放到刚好放下整张图"],
            ["键盘 ↑ ↓", "在节点之间切换选中"],
            ["键盘 ← →", "→ 展开 / 进入子节点，← 折叠 / 回到父节点"],
            ["Enter", "打开选中的节点"],
            ["0（数字零）/ + / -", "适应视图 / 放大 / 缩小"],
        ],
        widths=[5.6, 10.6],
    )
    body(doc, "思维导图只是**预览**：展开、拖动都不会修改磁盘上的文件。")

    # ------------------------------------------------------------------ 5
    h1(doc, "5. 快捷面板（HUD）")
    body(doc, "如果只是临时找一个网站或文件夹，不必打开主窗口——用快捷键直接弹出一个小面板，"
              "搜完就走。面板是一个 480×600 的悬浮窗口，按分类分组显示资源。")
    figure(doc, "quick-panel.png", "文件夹快捷面板（Tab 可切换到网站树）", width_cm=8.2)

    h2(doc, "5.1 默认快捷键")
    table(
        doc,
        ["快捷键", "作用"],
        [
            ["Ctrl + Alt + Space", "显示 / 隐藏主窗口"],
            ["Ctrl + Alt + Shift + L", "弹出「网站」快捷面板"],
            ["Ctrl + Alt + Shift + F", "弹出「文件夹」快捷面板"],
        ],
        widths=[5.8, 10.4],
    )

    h2(doc, "5.2 面板里的操作")
    bullet(doc, "直接输入即可过滤；面板会自动把焦点放在搜索框。")
    bullet(doc, "**↑ / ↓** 选择，**Enter** 打开，**Esc** 关闭面板。")
    bullet(doc, "文件夹面板里按 **→** 展开当前文件夹的子项；单击文件夹会直接打开资源管理器。")
    bullet(doc, "**Tab** 在网站树和文件夹树之间切换。")
    bullet(doc, "鼠标点击面板**外部**会自动关闭；再按一次打开它的快捷键也会关闭。")

    h2(doc, "5.3 自定义快捷键")
    body(doc, "打开设置 →「快捷键组合」，点「录制快捷键」后按下想要的组合即可。"
              "如果组合被其它软件占用，设置里会提示，换一个组合保存就行。"
              "三个快捷键（主窗口、网站面板、文件夹面板）分别设置，互不重复。")

    # ------------------------------------------------------------------ 6
    h1(doc, "6. 托盘、窗口与个性化")
    h2(doc, "6.1 关闭窗口 = 收进托盘")
    body(doc, "点窗口右上角的 × 不会退出程序，而是隐藏到系统托盘。此时：")
    bullet(doc, "**左键单击 / 双击**托盘图标：显示或隐藏主窗口。")
    bullet(doc, "**右键单击**托盘图标：菜单里可以「显示 / 隐藏」或「退出程序」。")
    bullet(doc, "全局快捷键 `Ctrl + Alt + Space` 也能随时把窗口叫出来。")

    h2(doc, "6.2 单实例")
    body(doc, "抽屉同时只能运行一个。重复双击 exe 不会打开第二个窗口，而是把已有窗口叫到前台。"
              "要彻底退出请用托盘菜单的「退出程序」。")

    h2(doc, "6.3 界面语言与外观")
    bullet(doc, "**界面语言**：设置 → 界面语言 → 简体中文 / English，立即生效并保存。")
    bullet(doc, "**外观主题**：设置 → 外观主题 → 浅色 / 跟随系统 / 深色，立即生效。"
                 "（主题是本次运行的显示偏好，重启后默认回到「跟随系统」。）")

    h2(doc, "6.4 开机启动")
    body(doc, "设置 → 开机启动：打开后会在当前用户的启动文件夹创建快捷方式，登录 Windows 后"
              "静默进入托盘；关闭则删除该快捷方式。")

    h2(doc, "6.5 设置窗口")
    body(doc, "所有开关和选项**改完立即生效**，不需要点保存；右上角的 × 关闭设置。"
              "书签目录、AI 密钥等需要输入的内容在输入框失焦或点按钮时保存。")

    # ------------------------------------------------------------------ 7
    h1(doc, "7. 浏览器书签同步（Chrome / Edge）")
    h2(doc, "7.1 自动导入：只读、不修改浏览器")
    body(doc, "抽屉会读取本机 Chrome / Edge 的书签文件，把书签自动变成「网站」资源："
              "书签文件夹名作为分类（书签栏、其他收藏等根目录也一样），并自动去重。"
              "启动时会同步一次，运行中每两秒检查一次，浏览器里新加的书签会很快出现。")
    bullet(doc, "抽屉**从不修改**浏览器书签文件，只是读取。")
    bullet(doc, "同一个网址只会自动导入一次；在工具里删除后，不会因为书签还在而反复加回来。")
    bullet(doc, "在浏览器里取消收藏，默认**不会**从抽屉里删除（需要下一节的双向同步扩展）。")

    h2(doc, "7.2 设置里怎么调")
    bullet(doc, "**启用书签同步**：关闭后不再自动导入。")
    bullet(doc, "**自定义书签目录**：默认「自动检测」；便携版浏览器、多用户配置或非标准安装位置，"
                 "可用「选择目录」指定 User Data 目录，或直接指定 `Bookmarks` 文件。")
    bullet(doc, "设置里会显示检测到的浏览器配置数量、书签总数、上次同步时间；"
                 "点「立即同步」可以手动触发一次。")
    bullet(doc, "**更多菜单 → 同步浏览器书签** 也能手动同步，结果显示在状态栏。")

    h2(doc, "7.3 双向同步（可选，装一个扩展）")
    body(doc, "如果希望「在抽屉里删网站 → 浏览器里的书签也一起取消」，需要加载随包的浏览器扩展"
              "（Chrome / Edge 通用，一次设置即可）：")
    numbered(
        doc,
        [
            "打开抽屉设置 →「浏览器书签同步」→ 点「打开扩展目录」，会打开程序里的 "
            "`bookmark-extension` 文件夹。",
            "在 Chrome 地址栏输入 `chrome://extensions`（Edge 用 `edge://extensions`）回车。",
            "打开右上角的「开发者模式」。",
            "点「加载已解压的扩展程序」，选择刚才打开的 `bookmark-extension` 文件夹。",
            "回到抽屉设置，看到「浏览器扩展：已连接」就成功了；如果显示未连接，重启一下浏览器，"
            "或到扩展页点一次刷新。",
        ],
    )
    body(doc, "装好之后：")
    bullet(doc, "删除同步来的网站时，可以勾选「同时从浏览器移除（N 处书签收藏）」，浏览器里的书签会被取消；"
                 "如果扩展暂时没运行，删除也会先记下，等扩展连上后自动补上。")
    bullet(doc, "在浏览器里删除或改名书签，抽屉会跟着同步（删除对应资源 / 更新名称）。")
    bullet(doc, "删除策略在设置里可选：每次询问 / 同时移除书签 / 仅删本地。")

    # ------------------------------------------------------------------ 8
    h1(doc, "8. 资源管理器右键菜单")
    body(doc, "抽屉会在资源管理器里加一个当前用户（不需要管理员权限）的右键菜单项，"
              "方便把任意文件夹直接收进抽屉：")
    bullet(doc, "在**文件夹图标**上点右键 →「添加到抽屉」。")
    bullet(doc, "在**文件夹里的空白处**点右键 →「添加当前文件夹到抽屉」。")
    body(doc, "点完之后抽屉会弹出，并自动选中新加入的文件夹。Windows 11 默认折叠旧版菜单，"
              "需要先点「显示更多选项」才能看到。")
    bullet(doc, "设置 →「资源管理器右键菜单」可以随时开关；关闭后菜单项会自动删除。")
    bullet(doc, "移动了抽屉的文件夹位置后，**重新启动一次抽屉**，菜单会自动指向新位置。")

    # ------------------------------------------------------------------ 9
    h1(doc, "9. AI 分类建议（实验功能）")
    body(doc, "这是可选功能，**默认关闭**。它会把「未分类」的网站和文件夹发给 AI，"
              "让 AI 建议一个分类；你确认之后才会真正写入，不会自动修改任何东西。")

    h2(doc, "9.1 配置")
    numbered(
        doc,
        [
            "打开设置 →「AI 整理（实验）」，打开开关。",
            "填「API 地址」和「模型」。默认是 DeepSeek（`https://api.deepseek.com` / "
            "`deepseek-chat`）；也支持 OpenAI 或任何 OpenAI 兼容接口，见下表。",
            "把 API Key 粘贴到「API Key」输入框，点「保存密钥」。显示「API Key 已配置」即可。",
        ],
    )
    table(
        doc,
        ["服务", "API 地址", "模型示例"],
        [
            ["DeepSeek", "https://api.deepseek.com", "deepseek-chat"],
            ["OpenAI", "https://api.openai.com/v1", "gpt-4o-mini"],
            ["本地 Ollama", "http://localhost:11434/v1", "本地模型名，如 qwen2.5:7b"],
        ],
        widths=[3.6, 6.6, 6.0],
    )
    body(doc, "程序会在地址后面自动拼接 `/chat/completions`。保存密钥后，密钥用 Windows "
              "DPAPI 按当前用户加密保存在 `.resourcehub-ai-key` 文件里，`settings.json` 中不会出现明文。")

    h2(doc, "9.2 使用")
    numbered(
        doc,
        [
            "把分类筛选切到「未分类」（没有未分类资源时菜单里会提示没有可整理的内容）。",
            "点 **更多菜单 → AI 整理建议**，等待几秒。",
            "弹出的审查框里按「网站 / 文件夹」分组列出建议，默认全部勾选；"
            "可以取消不想要的条目。",
            "点「应用所选」，状态栏会提示应用了几条。",
        ],
    )
    body(doc, "一次最多处理 40 条未分类资源；只读取名称、网址 / 路径名等必要信息，"
              "**不会上传文件内容**，文件夹也只发送名称和上级目录名，不发送完整本地路径。"
              "调用云端接口可能产生少量费用（DeepSeek 通常很便宜），具体以你的服务商为准。")

    h2(doc, "9.3 常见提示")
    table(
        doc,
        ["提示", "含义 / 处理"],
        [
            ["AI 整理未启用", "设置里的 AI 开关没打开。"],
            ["API Key 未配置", "Key 没保存成功，或换了电脑/账户（加密文件无法跨用户解密），重新保存一次。"],
            ["AI 请求失败", "网络不通、地址或模型填错、余额不足、或接口返回格式异常；检查后重试。"],
        ],
        widths=[4.4, 11.8],
    )

    # ------------------------------------------------------------------ 10
    h1(doc, "10. 设置项总览")
    body(doc, "设置窗口的选项都是即时生效的；下面这张截图展示了主要分组。")
    figure(doc, "settings.png", "设置窗口（向下滚动还有书签同步、AI、右键菜单与开机启动）", width_cm=10.2)
    table(
        doc,
        ["分组", "设置项", "说明", "默认"],
        [
            ["通用", "界面语言", "简体中文 / English", "简体中文"],
            ["通用", "外观主题", "浅色 / 跟随系统 / 深色", "跟随系统"],
            ["快捷键", "启用全局快捷键", "总开关", "开"],
            ["快捷键", "快捷键组合", "显示 / 隐藏主窗口，点「录制快捷键」修改", "Ctrl + Alt + Space"],
            ["快捷键", "网站树快捷键", "弹出网站快捷面板", "Ctrl + Alt + Shift + L"],
            ["快捷键", "文件夹树快捷键", "弹出文件夹快捷面板", "Ctrl + Alt + Shift + F"],
            ["书签", "启用书签同步", "自动导入 Chrome / Edge 新书签", "开"],
            ["书签", "自定义书签目录", "自动检测或手动指定", "自动检测"],
            ["书签", "删除同步资源时", "每次询问 / 同时移除书签 / 仅删本地", "每次询问"],
            ["书签", "浏览器扩展", "显示已连接 / 未连接，可打开扩展目录", "—"],
            ["AI", "AI 整理（实验）", "总开关", "关"],
            ["AI", "API 地址 / 模型 / API Key", "DeepSeek 或 OpenAI 兼容接口", "DeepSeek"],
            ["系统", "资源管理器右键菜单", "显示「添加到抽屉」", "开"],
            ["系统", "开机启动", "登录后静默进托盘", "关"],
        ],
        widths=[2.4, 4.2, 7.4, 2.6],
        font_size=8.6,
    )

    # ------------------------------------------------------------------ 11
    h1(doc, "11. 数据、备份与迁移")
    h2(doc, "11.1 数据存在哪")
    body(doc, "所有数据都在程序所在的文件夹里，不写注册表（右键菜单和开机启动除外），也不上传云端：")
    table(
        doc,
        ["文件", "内容"],
        [
            ["sites.json", "网站收藏"],
            ["folders.json", "文件和文件夹收藏"],
            ["settings.json", "语言、快捷键、书签同步、AI 等设置"],
            [".resourcehub-meta.json", "书签去重记录、排序版本等状态"],
            [".resourcehub-ai-key", "AI Key 的加密文件（配置过才有）"],
        ],
        widths=[4.6, 11.6],
    )
    body(doc, "每次修改都会自动保存，并保留一份同名的 `.bak` 备份文件。")

    h2(doc, "11.2 导出、导入与换电脑")
    bullet(doc, "日常备份：**更多菜单 → 导出** 得到 JSON 文件，随时可用 **更多菜单 → 导入** 恢复。")
    bullet(doc, "换电脑 / 给同学：**把整个程序文件夹拷走**即可，数据和设置一起带走。"
                 "AI Key 因为按 Windows 用户加密，换电脑后需要重新填一次。")
    bullet(doc, "升级版本：把新版程序文件覆盖到旧文件夹（保留自己的 json 数据），"
                 "或在新文件夹首次启动后把旧的 `sites.json` / `folders.json` / `settings.json` 拷进来。")

    h2(doc, "11.3 卸载")
    numbered(
        doc,
        [
            "先在设置里关闭「资源管理器右键菜单」和「开机启动」；",
            "用托盘菜单「退出程序」彻底退出；",
            "删除整个程序文件夹即可（数据也在里面，如需保留请先导出）。",
        ],
    )
    body(doc, "如果顺序反了、已经删了文件夹：重新解压运行一次，关掉这两个开关后再删即可清理干净。")

    # ------------------------------------------------------------------ 12
    h1(doc, "12. 常见问题")
    table(
        doc,
        ["问题", "处理办法"],
        [
            ["双击 exe 没反应 / 闪退",
             "确认是整包解压而不是只拷了 exe；确认系统是 64 位 Windows 10/11；"
             "看杀毒软件是否拦截了 `backend\\resource_hub_backend.exe`，加信任后重试。"],
            ["SmartScreen 提示或杀毒误报",
             "未签名程序的正常提示：点「更多信息 → 仍要运行」；杀软里把程序文件夹加入信任。"],
            ["数据保存不下来",
             "程序被放在 Program Files 等只读目录，移到 D 盘或桌面上的普通文件夹。"],
            ["快捷键没反应",
             "组合被其它软件占用：设置里重新录制一个；也确认「启用全局快捷键」是开着的。"],
            ["右键菜单没有「添加到抽屉」",
             "Windows 11 需要先点「显示更多选项」；确认设置里开关是开的；"
             "移动过程序文件夹的话，重启一次抽屉。"],
            ["浏览器书签没导入",
             "确认书签同步开关是开的；便携版浏览器用「自定义书签目录」指定；"
             "点「立即同步」；新书签最多等几秒。"],
            ["扩展显示「未连接」",
             "确认扩展已在扩展页加载并启用；重启浏览器；确认抽屉正在运行；"
             "Chrome / Edge 官方版需要在开发者模式下「加载已解压的扩展程序」。"],
            ["删了网站，浏览器书签还在",
             "双向同步需要装扩展；或在删除确认框勾选「同时从浏览器移除」，"
             "或把删除策略设为「同时移除书签」。"],
            ["AI 不能用",
             "确认 AI 开关已开、Key 已保存、地址和模型正确、网络可用；没有未分类资源时不会请求。"],
            ["资源显示路径失效",
             "用「检测路径」找出失效项，编辑资源更新路径，或删除后重新添加。"],
            ["重复双击 exe 没有新窗口",
             "抽屉是单实例程序，会唤出已经运行的窗口，这是正常的。"],
        ],
        widths=[5.2, 11.0],
        font_size=8.8,
    )

    # ------------------------------------------------------------------ 13
    h1(doc, "13. 快捷键与手势速查")
    table(
        doc,
        ["场景", "操作", "作用"],
        [
            ["任意位置", "Ctrl + Alt + Space", "显示 / 隐藏主窗口"],
            ["任意位置", "Ctrl + Alt + Shift + L", "网站快捷面板"],
            ["任意位置", "Ctrl + Alt + Shift + F", "文件夹快捷面板"],
            ["主窗口", "Alt + ↑ / Alt + ↓", "选中的资源上移 / 下移"],
            ["主窗口", "Esc", "关闭详情面板 / 对话框"],
            ["思维导图", "拖动", "平移视图"],
            ["思维导图", "滚轮 / Shift + 滚轮", "上下 / 左右平移"],
            ["思维导图", "Ctrl + 滚轮", "缩放"],
            ["思维导图", "双击空白", "适应视图"],
            ["思维导图", "方向键 / Enter", "选择、展开、打开节点"],
            ["思维导图", "0 / + / -", "适应 / 放大 / 缩小"],
            ["快捷面板", "↑ ↓ / → / Enter / Esc / Tab", "选择 / 展开 / 打开 / 关闭 / 切换树"],
        ],
        widths=[3.2, 6.6, 6.4],
        font_size=8.8,
    )

    # ------------------------------------------------------------------ 14
    h1(doc, "14. 隐私与安全")
    bullet(doc, "**本地优先**：收藏、设置、书签导入都在本机完成，不注册账号，不向任何服务器上传资源数据。")
    bullet(doc, "**联网只有一个可选用途**：AI 分类建议。请求只发给你自己填的 API 地址，内容仅限"
                 "未分类资源的名称、网址、分类和文件夹名称（不含完整路径和文件内容）。")
    bullet(doc, "**API Key**：用 Windows DPAPI 按当前用户加密保存在 `.resourcehub-ai-key`，"
                 "不写入 `settings.json`；换 Windows 账户或换电脑后无法解密，需要重新填写。")
    bullet(doc, "**浏览器扩展**：只申请 `bookmarks` 权限，通过本机回环地址（127.0.0.1）与抽屉通信，"
                 "不连接外网；不安装扩展时抽屉只读取书签，双向删除不可用。")
    bullet(doc, "**可随时退出**：删除程序文件夹即删除全部本地数据；注销右键菜单和开机启动见 11.3。")

    # ------------------------------------------------------------------ 15
    h1(doc, "15. 版本与反馈")
    body(doc, "当前版本：**v1.0（2026-10-18）**，Windows 10 / 11 64 位。")
    body(doc, "后续更新会以「整个程序文件夹」的形式发布：下载新版本 → 覆盖程序文件、保留自己的 "
              "json 数据（或直接用导出/导入迁移）。")
    body(doc, "问题反馈、功能建议或想看源码：")
    bullet(doc, "GitHub：`https://github.com/hsq2004815-oss/resource-hub-desktop`")
    bullet(doc, "Gitee：`https://gitee.com/hsqshuangquandyy/resource-hub-desktop`")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(18)
    add_rich(spacer, "感谢使用抽屉。希望它能帮你少翻几次文件夹。", size=10, color=MUTED)

    return doc


def main() -> int:
    doc = build()
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    size = OUT.stat().st_size
    print(f"saved {OUT} ({size / 1024:.1f} KB)")
    # Reopen for a cheap structural check.
    reopened = Document(str(OUT))
    headings = [p.text for p in reopened.paragraphs if p.style.name.startswith("Heading 1")]
    images = len(reopened.inline_shapes)
    tables = len(reopened.tables)
    print(f"headings={len(headings)} images={images} tables={tables}")
    print("h1:", headings)
    return 0


if __name__ == "__main__":
    sys.exit(main())
